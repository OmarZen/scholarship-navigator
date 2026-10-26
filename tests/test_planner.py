"""
Unit tests for Scholarship Finder & Application Planner.
"""

import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from planner.models import CandidateProfile, ApplicationState, DocumentRecord
from planner.matcher import load_scholarships, match_profile_all, evaluate_scholarship_match, to_ielts
from planner.documents import DocumentInventory, load_catalog, inventory_from_dicts
from planner.regions import normalize_country, in_group
from planner.roadmap import generate_roadmap, select_shortlist, build_action_plan
from planner.xlsx_generator import generate_application_workbook
from planner.calendar_exporter import export_calendar_ics
from planner.essay_coach import review_scholarship_essay
from planner.outreach_generator import generate_supervisor_email
from planner.financial_calculator import calculate_financial_gap
from planner.folder_organizer import initialize_application_workspace
from planner.html_generator import generate_html_report
from planner.portfolio_strategist import analyze_application_portfolio
from planner.lor_generator import generate_recommendation_letter, generate_recommender_brag_sheet
from planner.credential_guide import get_credential_guide_for_country, load_credential_guides
from planner.interview_prep import get_interview_prep
from planner.fee_waivers import load_fee_waiver_data, generate_fee_waiver_email
from planner.relocation_guide import get_relocation_guide
from planner import workspace as wsp

TODAY = date(2026, 10, 25)
ROOT = Path(__file__).parent.parent


def by_id(matches, sid):
    return next(m for m in matches if m.scholarship.id == sid)


class TestScholarshipPlanner(unittest.TestCase):
    def setUp(self):
        self.profile = CandidateProfile(
            full_name="Sarah Ibrahim",
            nationality="Egypt",
            current_degree="Bachelor of Science in Computer Engineering",
            target_degree="Master",
            field_of_study="Artificial Intelligence",
            gpa=3.8,
            gpa_scale=4.0,
            english_test="IELTS Academic",
            english_score=8.0,
            work_experience_years=3.0,
            research_publications=1,
            target_funding="Full",
            target_countries=["Germany", "United Kingdom", "Hungary"],
            target_intake_year=2027,
            budget_for_applications="Moderate",
            has_leadership_experience=True,
            date_of_birth=date(1999, 6, 1),
            graduation_year=2023,
        )
        self.tmp = tempfile.TemporaryDirectory()
        self.tmpdir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def matches(self, profile=None, **kw):
        return match_profile_all(profile or self.profile, today=TODAY, **kw)

    # -- data integrity -------------------------------------------------------

    def test_load_scholarships(self):
        scholarships = load_scholarships()
        self.assertGreaterEqual(len(scholarships), 30)
        ids = [s.id for s in scholarships]
        for sid in ("chevening", "erasmus-mundus", "stipendium-hungaricum", "daad-epos", "fulbright", "mbzuai-ai", "vanier-cgs"):
            self.assertIn(sid, ids)

    def test_every_requirement_maps_to_catalog(self):
        catalog = load_catalog()
        for s in load_scholarships():
            self.assertTrue(s.documents, f"{s.id} has no documents")
            # Unknown stays unknown: no pattern is allowed only when the published text explains the deadline.
            self.assertTrue(s.timeline.deadline or s.timeline.typical_deadline, f"{s.id} has no deadline information")
            for req in s.documents:
                self.assertIn(req.doc, catalog, f"{s.id}: unknown doc id {req.doc}")
                for alt in req.alternatives:
                    self.assertIn(alt, catalog, f"{s.id}: unknown alternative {alt}")

    def test_skill_copies_in_sync(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "sync_skills.py"), "--check"],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, "Run python scripts/sync_skills.py\n" + result.stdout)
        for skill in (ROOT / "skills").iterdir():
            text = (skill / "SKILL.md").read_text(encoding="utf-8")
            self.assertTrue(text.startswith("---\nname: " + skill.name + "\n"), f"{skill.name}: bad front matter")
            self.assertIn("description:", text.split("---", 2)[1])

    def test_catalog_docs_up_to_date(self):
        result = subprocess.run([sys.executable, str(ROOT / "scripts" / "build_catalog_docs.py"), "--check"],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, "Run python scripts/build_catalog_docs.py\n" + result.stdout)

    def test_readme_links_exist(self):
        import re
        for doc in ("README.md", "docs/i18n/README.ar.md", "CONTRIBUTING.md", "docs/MAINTAINING.md"):
            path = ROOT / doc
            text = path.read_text(encoding="utf-8")
            targets = re.findall(r"\]\(([^)#\s]+)", text) + re.findall(r'(?:src|href)="([^"#]+)"', text)
            for target in targets:
                if target.startswith(("http://", "https://", "mailto:")):
                    continue
                self.assertTrue((path.parent / target).exists(), f"{doc}: broken link {target}")
        for required in ("LICENSE", "CODE_OF_CONDUCT.md", "SECURITY.md", "CHANGELOG.md", ".github/CODEOWNERS",
                         ".github/rulesets/protect-main.json", ".github/pull_request_template.md"):
            self.assertTrue((ROOT / required).exists(), required)

    def test_data_validates(self):
        from planner.validation import validate_data
        res = validate_data(today=TODAY)
        self.assertTrue(res["ok"], res["errors"])
        self.assertGreaterEqual(res["counts"]["programs"], 50)

    def test_schemas_up_to_date(self):
        from planner.validation import schema_documents
        for name, text in schema_documents().items():
            path = ROOT / "schemas" / f"{name}.schema.json"
            self.assertTrue(path.exists(), f"missing {path}: run cli.py export-schema")
            self.assertEqual(path.read_text(encoding="utf-8"), text, f"{path.name} is stale: run cli.py export-schema")

    def test_country_normalization(self):
        self.assertEqual(normalize_country("Egyptian"), "Egypt")
        self.assertEqual(normalize_country("UK"), "United Kingdom")
        self.assertTrue(in_group("Cape Verde", "dac_oda"))
        self.assertTrue(in_group("egypt", "africa"))
        self.assertFalse(in_group("Germany", "dac_oda"))
        self.assertIsNone(in_group("Turkey", "europe"))

    def test_language_concordance(self):
        self.assertEqual(to_ielts("IELTS Academic", 7.0), 7.0)
        self.assertEqual(to_ielts("TOEFL iBT", 100), 7.0)
        self.assertEqual(to_ielts("Duolingo", 125), 6.5)
        self.assertIsNone(to_ielts(None, None))

    # -- eligibility gates & scoring -----------------------------------------

    def test_matching_algorithm(self):
        matches = self.matches()
        top = matches[0]
        self.assertEqual(top.match_status, "High Match")
        self.assertGreaterEqual(top.match_score, 80.0)
        self.assertTrue(top.score_breakdown)
        self.assertTrue(top.gate_checks)
        self.assertGreater(sum(c.weight for c in top.score_breakdown), 0)

    def test_phd_disqualification_for_master_profile(self):
        singa = by_id(self.matches(), "singa-singapore")
        self.assertEqual(singa.match_status, "Ineligible")
        self.assertEqual(singa.match_score, 0.0)

    def test_nationality_gate(self):
        german = self.profile.model_copy(update={"nationality": "Germany"})
        mastercard = by_id(self.matches(german), "mastercard-foundation")
        self.assertEqual(mastercard.eligibility_status, "Ineligible")
        self.assertTrue(any(g.name == "Nationality" and g.result == "fail" for g in mastercard.gate_checks))
        egypt = by_id(self.matches(), "mastercard-foundation")
        self.assertNotEqual(egypt.eligibility_status, "Ineligible")

    def test_missing_birth_date_needs_check(self):
        no_dob = self.profile.model_copy(update={"date_of_birth": None})
        csc = by_id(self.matches(no_dob), "csc-china")
        self.assertEqual(csc.eligibility_status, "Needs Check")
        self.assertTrue(any(g.name == "Age" and g.result == "unknown" and g.fix for g in csc.gate_checks))

    def test_age_limit(self):
        older = self.profile.model_copy(update={"date_of_birth": date(1990, 1, 1)})
        turkey = by_id(self.matches(older), "turkiye-burslari")
        self.assertEqual(turkey.eligibility_status, "Ineligible")

    def test_work_experience_gate(self):
        junior = self.profile.model_copy(update={"work_experience_years": 0.5, "currently_employed": False})
        chevening = by_id(self.matches(junior), "chevening")
        self.assertEqual(chevening.eligibility_status, "Ineligible")
        self.assertTrue(any(g.name == "Work experience" and g.result == "fail" for g in chevening.gate_checks))

    def test_restricted_field(self):
        engineer = by_id(self.matches(), "daad-helmut-schmidt")
        self.assertTrue(any(g.name == "Field of study" and g.result == "fail" for g in engineer.gate_checks))
        policy = self.profile.model_copy(update={"field_of_study": "Public Policy and Governance",
                                                 "current_degree": "Bachelor of Political Science"})
        hs = by_id(self.matches(policy), "daad-helmut-schmidt")
        self.assertTrue(any(g.name == "Field of study" and g.result == "pass" for g in hs.gate_checks))

    def test_discontinued_program_is_ineligible(self):
        phd = self.profile.model_copy(update={"target_degree": "PhD"})
        vanier = by_id(self.matches(phd), "vanier-cgs")
        self.assertEqual(vanier.eligibility_status, "Ineligible")
        self.assertTrue(any(g.name == "Program status" for g in vanier.gate_checks))

    # -- study location (home / abroad / both), gender, diploma ---------------

    def gate(self, m, name):
        return next((g for g in m.gate_checks if g.name == name), None)

    def test_study_location_home_and_abroad(self):
        saudi = self.profile.model_copy(update={"nationality": "Saudi Arabia", "target_countries": []})
        home = self.matches(saudi.model_copy(update={"study_location": "home"}))
        kaust = by_id(home, "kaust-fellowship")
        self.assertEqual(self.gate(kaust, "Study location").result, "pass")
        self.assertTrue(kaust.study_at_home)
        for sid in ("chevening", "erasmus-mundus", "mbzuai-ai"):
            m = by_id(home, sid)
            self.assertEqual(m.eligibility_status, "Ineligible", sid)
            self.assertEqual(self.gate(m, "Study location").result, "fail", sid)

        abroad = self.matches(saudi.model_copy(update={"study_location": "abroad"}))
        self.assertEqual(self.gate(by_id(abroad, "kaust-fellowship"), "Study location").result, "fail")
        self.assertEqual(self.gate(by_id(abroad, "mbzuai-ai"), "Study location").result, "pass")

        both = self.matches(saudi)
        self.assertEqual(saudi.study_location, "both")
        self.assertIsNone(self.gate(by_id(both, "kaust-fellowship"), "Study location"))

    def test_region_hosts_and_mobility_schemes(self):
        kenyan = self.profile.model_copy(update={"nationality": "Kenya", "study_location": "home", "target_countries": []})
        matches = self.matches(kenyan)
        mastercard = by_id(matches, "mastercard-foundation")
        self.assertEqual(self.gate(mastercard, "Study location").result, "unknown")
        self.assertIsNone(mastercard.study_at_home)
        self.assertEqual(self.gate(by_id(matches, "intra-africa-mobility"), "Study location").result, "fail")

    def test_both_accepts_home_country_outside_targets(self):
        saudi = self.profile.model_copy(update={"nationality": "Saudi Arabia", "target_countries": ["Germany"]})
        prefs = {c.key: c for c in by_id(self.matches(saudi), "kaust-fellowship").score_breakdown}["preferences"]
        self.assertIn("at home", prefs.detail)
        abroad_only = saudi.model_copy(update={"study_location": "abroad"})
        prefs = {c.key: c for c in by_id(self.matches(abroad_only), "kaust-fellowship").score_breakdown}["preferences"]
        self.assertIn("outside your target countries", prefs.detail)

    def test_gender_gate(self):
        s = next(x for x in load_scholarships() if x.id == "erasmus-mundus")
        rules = s.rules.model_copy(update={"gender": "female"})
        women_only = s.model_copy(update={"eligibility": s.eligibility.model_copy(update={"rules": rules})})
        results = {g: evaluate_scholarship_match(self.profile.model_copy(update={"gender": g}), women_only, today=TODAY)
                   for g in ("female", "male", None)}
        self.assertEqual(self.gate(results["female"], "Gender").result, "pass")
        self.assertEqual(results["male"].eligibility_status, "Ineligible")
        self.assertEqual(results[None].eligibility_status, "Needs Check")
        self.assertIsNone(self.gate(evaluate_scholarship_match(self.profile, s, today=TODAY), "Gender"))

    def test_diploma_degree_level(self):
        s = next(x for x in load_scholarships() if x.id == "kaust-fellowship")
        diploma_award = s.model_copy(update={"degrees": ["Diploma"]})
        student = self.profile.model_copy(update={"target_degree": "Diploma", "work_experience_years": 0})
        m = evaluate_scholarship_match(student, diploma_award, today=TODAY)
        self.assertEqual(self.gate(m, "Degree level").result, "pass")
        self.assertNotIn("experience", {c.key for c in m.score_breakdown})
        self.assertEqual(evaluate_scholarship_match(student, s, today=TODAY).eligibility_status, "Ineligible")

    def test_list_scholarships_by_nationality_and_location(self):
        from planner import api
        home = {r["id"]: r for r in api.list_scholarships(nationality="Saudi Arabia", study_location="home")}
        self.assertIn("kaust-fellowship", home)
        self.assertNotIn("chevening", home)
        self.assertEqual(home["kaust-fellowship"]["study_location"], "Home country")
        abroad = {r["id"] for r in api.list_scholarships(nationality="Saudi Arabia", study_location="abroad")}
        self.assertNotIn("kaust-fellowship", abroad)
        self.assertIn("mbzuai-ai", abroad)
        self.assertNotIn("singa-singapore", {r["id"] for r in api.list_scholarships(nationality="Singapore")})
        with self.assertRaises(ValueError):
            api.list_scholarships(study_location="home")

    def test_programs_respect_study_location(self):
        from planner.programs import rank_programs
        home = self.profile.model_copy(update={"study_location": "home"})
        self.assertFalse(rank_programs(home, today=TODAY))
        self.assertTrue(rank_programs(self.profile.model_copy(update={"study_location": "abroad"}), today=TODAY))

    def test_set_study_location_in_workspace(self):
        ws, _ = self._workspace()
        wsp.set_study_location(ws, "home")
        raw = json.loads(wsp.profile_path(ws).read_text(encoding="utf-8"))
        self.assertEqual(raw["study_location"], "home")
        self.assertIn("_help", raw)
        res = wsp.build_workspace(ws, today=TODAY)
        self.assertEqual(res["profile"].study_location, "home")
        self.assertIn("Home country only", (ws / "PROFILE.md").read_text(encoding="utf-8"))
        self.assertTrue(all(m.study_at_home is not False for m in res["shortlist"]))
        with self.assertRaises(ValueError):
            wsp.set_study_location(ws, "moon")

    # -- readiness & deadlines ------------------------------------------------

    def test_readiness_counts_documents(self):
        s = next(x for x in load_scholarships() if x.id == "erasmus-mundus")
        empty = evaluate_scholarship_match(self.profile, s, inventory=DocumentInventory(), today=TODAY)
        inv = inventory_from_dicts([
            {"doc_id": "passport", "status": "have"},
            {"doc_id": "degree_certificate", "status": "have"},
            {"doc_id": "lor_academic", "status": "have", "count_ready": 1},
        ])
        some = evaluate_scholarship_match(self.profile, s, inventory=inv, today=TODAY)
        self.assertGreater(some.readiness_percent, empty.readiness_percent)
        lor = next(r for r in some.readiness.requirements if r.doc == "lor_academic")
        self.assertEqual(lor.state, "partial")
        self.assertAlmostEqual(lor.units, 0.5)
        # English test reported in the profile counts as held.
        eng = next(r for r in empty.readiness.requirements if r.doc == "english_test")
        self.assertEqual(eng.state, "have")

    def test_expiring_document(self):
        s = next(x for x in load_scholarships() if x.id == "swedish-institute")
        inv = inventory_from_dicts([{"doc_id": "english_test", "status": "have", "issued": "2024-11-01"}])
        m = evaluate_scholarship_match(self.profile, s, inventory=inv, today=TODAY)
        eng = next(r for r in m.readiness.requirements if r.doc == "english_test")
        self.assertEqual(eng.state, "expiring")
        self.assertLessEqual(eng.units, 0.5)

    def test_deadline_passed_and_next_cycle(self):
        chevening = by_id(self.matches(), "chevening")
        self.assertEqual(chevening.readiness.timeline_status, "Deadline passed")
        self.assertIsNotNone(chevening.readiness.next_cycle_deadline)
        self.assertGreater(chevening.readiness.next_cycle_deadline, TODAY)
        shortlist = select_shortlist(self.matches())
        self.assertNotIn("chevening", [m.scholarship.id for m in shortlist])

    def test_confirmed_deadline_overrides_estimate(self):
        state = ApplicationState(scholarship_id="erasmus-mundus", deadline_confirmed=date(2027, 1, 3))
        m = by_id(self.matches(states={"erasmus-mundus": state}), "erasmus-mundus")
        self.assertEqual(m.readiness.deadline, date(2027, 1, 3))
        self.assertFalse(m.readiness.deadline_is_estimate)

    def test_roadmap_generation(self):
        roadmap = generate_roadmap(self.profile, self.matches(), today=TODAY)
        self.assertEqual(roadmap["candidate"], "Sarah Ibrahim")
        self.assertEqual(len(roadmap["phases"]), 5)
        self.assertTrue(roadmap["actions"])

    def test_action_plan_is_dated_and_sorted(self):
        shortlist = select_shortlist(self.matches())
        actions = build_action_plan(self.profile, shortlist, TODAY)
        dated = [a for a in actions if a["bucket"] != "After submitting"]
        self.assertTrue(all(a["start_by"] for a in dated))
        self.assertTrue(any(a["kind"] == "submit" for a in actions))
        order = ["Start now", "Next 30 days", "1-3 months", "3+ months", "After submitting"]
        buckets = [order.index(a["bucket"]) for a in actions]
        self.assertEqual(buckets, sorted(buckets))

    # -- workspace ------------------------------------------------------------

    def _workspace(self):
        data = self.profile.model_dump(mode="json")
        data["documents"] = [{"doc_id": "passport", "status": "have", "expires": "2031-01-01"}]
        ws = wsp.create_workspace(data, root=self.tmpdir)
        return ws, wsp.build_workspace(ws, today=TODAY)

    def test_workspace_structure(self):
        ws, res = self._workspace()
        for rel in ("AGENTS.md", "INDEX.md", "PROFILE.md", "manifest.json", "profile/profile.json",
                    "profile/documents.json", "profile/story_bank.md", "profile/identity/README.md",
                    "reports/scholarship_tracker.xlsx", "reports/scholarship_report.html", "reports/deadlines.ics",
                    "reports/match_report.md", "reports/action_plan.md"):
            self.assertTrue((ws / rel).exists(), rel)
        self.assertTrue(res["shortlist"])
        for m in res["shortlist"]:
            folder = ws / "scholarships" / m.scholarship.id
            self.assertTrue((folder / "SCHOLARSHIP.md").exists())
            self.assertTrue((folder / "status.json").exists())
        manifest = json.loads((ws / "manifest.json").read_text())
        self.assertEqual(manifest["read_first"][0], "AGENTS.md")
        self.assertEqual(len(manifest["shortlist"]), len(res["shortlist"]))

    def test_rebuild_preserves_user_drafts(self):
        ws, res = self._workspace()
        sid = res["shortlist"][0].scholarship.id
        drafts = [p for p in (ws / "scholarships" / sid).rglob("*.md") if p.name != "SCHOLARSHIP.md"]
        self.assertTrue(drafts)
        drafts[0].write_text(drafts[0].read_text().replace("TODO: write here.", "My real draft."), encoding="utf-8")
        wsp.build_workspace(ws, today=TODAY)
        self.assertIn("My real draft.", drafts[0].read_text())

    def test_final_draft_marks_paper_done(self):
        ws, res = self._workspace()
        target = next((m for m in res["shortlist"]
                       if any((ws / "scholarships" / m.scholarship.id / "letters").glob("*.md"))), None)
        if target is None:
            self.skipTest("no letter drafts in shortlist")
        letters = sorted((ws / "scholarships" / target.scholarship.id / "letters").glob("*.md"))
        text = letters[0].read_text()
        letters[0].write_text(text.replace("status: todo", "status: final"), encoding="utf-8")
        wsp.build_workspace(ws, today=TODAY)
        state = wsp.load_states(ws)[target.scholarship.id]
        doc = text.split("doc: ", 1)[1].split("\n", 1)[0]
        self.assertEqual(state.documents[doc].status, "have")

    def test_set_document_and_deadline(self):
        ws, _ = self._workspace()
        wsp.set_document(ws, "transcript", status="have", file="profile/academic/transcript.pdf", translated="yes")
        rec = wsp.load_inventory(ws).get("transcript")
        self.assertEqual(rec.status, "have")
        self.assertTrue(rec.translated)
        wsp.set_deadline(ws, "erasmus-mundus", "2027-01-05", "https://example.org", today=TODAY)
        self.assertEqual(wsp.load_states(ws)["erasmus-mundus"].deadline_confirmed, date(2027, 1, 5))

    def test_workbook_sync_round_trip(self):
        ws, _ = self._workspace()
        path = ws / "reports" / "scholarship_tracker.xlsx"
        wb = openpyxl.load_workbook(path)
        for row in wb["My Documents"].iter_rows(min_row=2):
            if row[0].value == "transcript":
                row[3].value = "have"
        wb.save(path)
        counts = wsp.apply_workbook_edits(ws)
        self.assertEqual(counts["documents"], 1)
        self.assertEqual(wsp.load_inventory(ws).get("transcript").status, "have")

    def test_legacy_folder_organizer(self):
        path = initialize_application_workspace(self.profile, None, str(self.tmpdir / "legacy_ws"))
        self.assertTrue((Path(path) / "INDEX.md").exists())


    # -- programs (Erasmus Mundus masters, Intra-Africa projects) -------------

    def test_programs_load_and_have_hosts(self):
        from planner.programs import load_programs
        programs = load_programs()
        self.assertGreaterEqual(len(programs), 50)
        ids = {p.id for p in programs}
        for pid in ("codas", "ediss", "steps", "ipcvai", "cybersure", "sustagri"):
            self.assertIn(pid, ids)
        emjm = [p for p in programs if p.parent_scholarship == "erasmus-mundus"]
        self.assertTrue(all(p.host_countries for p in emjm if p.status != "discontinued"))

    def test_program_ranking_by_field(self):
        from planner.programs import rank_programs
        ai = self.profile.model_copy(update={"field_of_study": "Artificial Intelligence and machine learning"})
        top = [pm.program.id for pm in rank_programs(ai, today=TODAY)[:8]]
        self.assertTrue({"ediss", "deai", "multicom", "ipcvai", "aiss"} & set(top), top)
        agro = self.profile.model_copy(update={"field_of_study": "Agronomy and food security",
                                               "current_degree": "BSc Agriculture", "target_countries": []})
        ranked = rank_programs(agro, today=TODAY)
        self.assertEqual(ranked[0].program.id, "sustagri")
        cyber = next(pm for pm in ranked if pm.program.id == "cybersure")
        self.assertEqual(cyber.field_match, "none")  # "food security" must not match cybersecurity

    def test_discontinued_and_africa_only_programs(self):
        from planner.programs import load_programs, match_program
        programs = {p.id: p for p in load_programs()}
        discontinued = [p for p in programs.values() if p.status == "discontinued"]
        self.assertTrue(discontinued)
        self.assertFalse(match_program(self.profile, discontinued[0], TODAY).eligible)
        spree = programs["spree"]
        german = self.profile.model_copy(update={"nationality": "Germany"})
        self.assertFalse(match_program(german, spree, TODAY).eligible)
        self.assertTrue(match_program(self.profile, spree, TODAY).eligible)  # Egypt is in Africa

    def test_chosen_programs_drive_erasmus_deadline(self):
        from planner.programs import load_programs, program_deadline
        programs = {p.id: p for p in load_programs()}
        choice = next(pid for pid, p in programs.items() if p.parent_scholarship == "erasmus-mundus"
                      and program_deadline(p, 2027)[0] and program_deadline(p, 2027)[0] > TODAY)
        state = ApplicationState(scholarship_id="erasmus-mundus", chosen_programs=[choice])
        m = by_id(self.matches(states={"erasmus-mundus": state}), "erasmus-mundus")
        self.assertEqual(m.readiness.deadline, program_deadline(programs[choice], 2027)[0])
        self.assertIn(programs[choice].acronym, m.readiness.deadline_label)

    def test_workspace_program_choice(self):
        ws, res = self._workspace()
        self.assertTrue(res["programs"])
        erasmus = ws / "scholarships" / "erasmus-mundus" / "SCHOLARSHIP.md"
        if erasmus.exists():
            self.assertIn("Programs that fit you", erasmus.read_text())
        pid = res["programs"][0].program.id
        wsp.choose_program(ws, pid)
        wsp.build_workspace(ws, today=TODAY)
        scheme = res["programs"][0].program.parent_scholarship
        self.assertIn(pid, wsp.load_states(ws)[scheme].chosen_programs)
        self.assertTrue((ws / "scholarships" / scheme / "letters" / f"motivation_letter_{pid}.md").exists())
        self.assertIn("Programs", openpyxl.load_workbook(ws / "reports" / "scholarship_tracker.xlsx").sheetnames)

    # -- agent interfaces -----------------------------------------------------

    def test_api_match_and_programs(self):
        from planner import api
        data = self.profile.model_dump(mode="json")
        data["documents"] = [{"doc_id": "passport", "status": "have"}]
        res = api.match(data, today=TODAY)
        self.assertTrue(res["results"])
        first = res["results"][0]
        for key in ("eligibility", "fit", "readiness_percent", "eligibility_checks", "fit_breakdown", "missing_papers"):
            self.assertIn(key, first)
        json.dumps(res, default=str)
        progs = api.find_programs(field="robotics", today=TODAY)
        self.assertTrue(all("robot" in " ".join([p["name"]] + p["matched_terms"]).lower() or p["field_match"] != "none"
                            for p in progs))

    def test_cli_json_output(self):
        out = subprocess.run([sys.executable, str(ROOT / "cli.py"), "match", "--profile",
                              str(ROOT / "examples" / "sample_candidate_full.json"), "--json", "--limit", "3"],
                             capture_output=True, text=True, cwd=ROOT)
        self.assertEqual(out.returncode, 0, out.stderr)
        parsed = json.loads(out.stdout)
        self.assertEqual(len(parsed), 3)
        self.assertIn("fit_breakdown", parsed[0])

    def test_mcp_server_tools(self):
        try:
            import mcp_server
        except ImportError:
            self.skipTest("mcp not installed")
        import asyncio
        names = {t.name for t in asyncio.run(mcp_server.mcp.list_tools())}
        for tool in ("match_scholarships", "find_programs", "create_workspace", "update_document",
                     "confirm_deadline", "choose_program", "list_document_types", "set_study_location"):
            self.assertIn(tool, names)
        result = mcp_server.find_programs(field="photonics", limit=5)
        self.assertTrue(result)

    def test_install_skills_script(self):
        script = str(ROOT / "scripts" / "install_skills.py")
        out = subprocess.run([sys.executable, script, "--tool", "codex", "--dest", str(self.tmpdir / "skills")],
                             capture_output=True, text=True)
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn("codex mcp add scholarship-navigator", out.stdout)
        installed = self.tmpdir / "skills" / "scholarship-intake" / "SKILL.md"
        text = installed.read_text(encoding="utf-8")
        self.assertTrue(text.startswith("---\nname: scholarship-intake\n"))
        self.assertIn("Installed copy", text)
        out = subprocess.run([sys.executable, script, "--tool", "claude-app", "--dest", str(self.tmpdir / "zips")],
                             capture_output=True, text=True)
        self.assertEqual(out.returncode, 0, out.stderr)
        import zipfile
        with zipfile.ZipFile(self.tmpdir / "zips" / "scholarship-finder-planner.zip") as z:
            self.assertIn("scholarship-finder-planner/SKILL.md", z.namelist())

    def test_contributors_block(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("update_contributors", ROOT / "scripts" / "update_contributors.py")
        uc = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(uc)
        people = [{"login": "alice", "type": "User", "html_url": "https://github.com/alice",
                   "avatar_url": "https://avatars.githubusercontent.com/u/1?v=4", "contributions": 3},
                  {"login": "dependabot[bot]", "type": "Bot", "html_url": "", "avatar_url": "", "contributions": 9}]
        block = uc.render(people)
        self.assertIn("@alice", block)
        self.assertNotIn("dependabot", block)
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn(block, uc.replace_block(readme, block))

    def test_custom_scholarship_scaffold_validates(self):
        from planner.validation import scholarship_template, validate_data
        path = self.tmpdir / "custom.json"
        path.write_text(json.dumps(scholarship_template("my-test-award")), encoding="utf-8")
        res = validate_data(today=TODAY, extra_files=[path])
        self.assertTrue(res["ok"], res["errors"])

    # -- reports --------------------------------------------------------------

    def test_xlsx_generation(self):
        out = self.tmpdir / "test_output.xlsx"
        generate_application_workbook(self.profile, self.matches(), str(out), today=TODAY)
        wb = openpyxl.load_workbook(out)
        for sheet in ("How to Use", "Dashboard", "Ranked Matches", "Requirements", "My Documents", "Action Plan",
                      "Timeline", "All Scholarships", "Country Guides"):
            self.assertIn(sheet, wb.sheetnames)
        readiness_col = [c.value for c in wb["Ranked Matches"][1]].index("Readiness")
        first = wb["Ranked Matches"].cell(row=2, column=readiness_col + 1).value
        self.assertTrue(str(first).startswith("=IFERROR(AVERAGEIFS("))

    def test_calendar_exporter(self):
        out = self.tmpdir / "test_calendar.ics"
        path = export_calendar_ics(self.profile, self.matches(), str(out), today=TODAY)
        content = Path(path).read_bytes().decode("utf-8")  # keep CRLF line endings intact
        self.assertIn("BEGIN:VCALENDAR", content)
        self.assertIn("BEGIN:VALARM", content)
        self.assertTrue(all(len(line.encode("utf-8")) <= 75 for line in content.split("\r\n")))

    def test_html_report_generation(self):
        out = self.tmpdir / "test_report.html"
        path = generate_html_report(self.profile, self.matches(), str(out), today=TODAY)
        content = Path(path).read_text(encoding="utf-8")
        self.assertIn("<!DOCTYPE html>", content)
        self.assertIn("Sarah Ibrahim", content)

    def test_portfolio_strategist(self):
        portfolio = analyze_application_portfolio(self.profile, self.matches())
        for key in ("reach_opportunities", "target_opportunities", "safety_opportunities",
                    "estimated_overall_success_probability", "portfolio_health", "probability_note"):
            self.assertIn(key, portfolio)

    # -- writing helpers & guides (unchanged modules) --------------------------

    def test_essay_reviewer(self):
        sample_text = (
            "When our engineering department faced a major server outage, I was tasked with restoring critical "
            "infrastructure within 48 hours. I spearheaded the migration to Kubernetes, negotiated with vendors, "
            "and organized 15 engineers into agile shifts. As a result, uptime was restored 12 hours ahead of schedule "
            "and configuration errors decreased by 45%."
        )
        res = review_scholarship_essay(sample_text, essay_type="chevening_leadership")
        self.assertGreaterEqual(res["readiness_score"], 60.0)
        self.assertTrue(len(res["strengths"]) > 0)

    def test_outreach_generator(self):
        email_data = generate_supervisor_email(self.profile, professor_name="Prof. John Doe", university_name="Oxford",
                                               laboratory_or_field="Machine Learning", relevant_paper_or_topic="Graph Neural Networks")
        self.assertIn("Prof. John Doe", email_data["body"])
        self.assertIn("Graph Neural Networks", email_data["body"])

    def test_financial_calculator(self):
        gap = calculate_financial_gap(country="Germany", scholarship_covers_tuition=True, monthly_stipend_received=992.0,
                                      stipend_duration_months=12, scholarship_covers_flights=True, scholarship_covers_insurance=True)
        self.assertEqual(gap["net_out_of_pocket_deficit"], 0.0)
        self.assertTrue(gap["is_fully_funded"])

    def test_lor_generator(self):
        lor = generate_recommendation_letter(profile=self.profile, recommender_name="Prof. Ahmed Hassan",
                                             recommender_title="Chair of Computer Engineering", institution_or_company="Cairo University",
                                             perspective="academic_research", relationship_years=3, scholarship_target="Chevening")
        self.assertIn("Prof. Ahmed Hassan", lor["body"])
        self.assertIn("Sarah Ibrahim", lor["body"])
        brag = generate_recommender_brag_sheet(self.profile, ["Chevening", "DAAD"])
        self.assertIn("CANDIDATE BRAG SHEET", brag)

    def test_credential_guide(self):
        self.assertGreaterEqual(len(load_credential_guides()), 4)
        self.assertIn("Uni-Assist", get_credential_guide_for_country("Germany")["evaluation_system"])

    def test_interview_prep(self):
        prep = get_interview_prep("chevening")
        self.assertGreaterEqual(len(prep["questions"]), 4)

    def test_fee_waivers(self):
        self.assertGreaterEqual(len(load_fee_waiver_data()), 2)
        email = generate_fee_waiver_email(profile=self.profile, university_name="TU Munich", program_name="MSc Informatics",
                                          department_chair_or_admissions="Admissions Committee")
        self.assertIn("TU Munich", email["body"])

    def test_relocation_guide(self):
        guide = get_relocation_guide("Germany")
        self.assertTrue(guide["visa_documents"] and guide["housing_platforms"] and guide["arrival_checklist"])


if __name__ == "__main__":
    unittest.main()
